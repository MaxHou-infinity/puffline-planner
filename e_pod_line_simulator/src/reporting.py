"""
报告导出模块 - 将仿真结果导出为 Excel / PDF

功能：
- export_excel(): 导出多 Sheet 的 Excel 报告（KPI、工序产出、报警、WIP 采样、切换事件）
- export_pdf(): 导出 PDF 摘要报告（内嵌中文字体）
- export_report(): 按文件扩展名自动分派
"""

import os
import json
import re
from typing import List, Optional
from xml.sax.saxutils import escape

from src.models import SimulationResult


def _keep_section_headers(story):
    """Keep a section heading and its spacer with the first following content."""
    for index, item in enumerate(story):
        if re.match(r"^[一二三四五六七八九十]+、", getattr(item, 'text', '')):
            item.keepWithNext = True
            if index + 1 < len(story) and story[index + 1].__class__.__name__ == 'Spacer':
                story[index + 1].keepWithNext = True


# 中文字体候选路径（按优先级）
_BUNDLED_CN_FONT = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "assets", "fonts", "NotoSansSC-Regular.ttf",
)
_CN_FONT_CANDIDATES = [
    _BUNDLED_CN_FONT,
    os.path.join(os.environ.get("WINDIR", "C:/Windows"), "Fonts", "msyh.ttc"),
    os.path.join(os.environ.get("WINDIR", "C:/Windows"), "Fonts", "simsun.ttc"),
    "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    "/System/Library/Fonts/Hiragino Sans GB.ttc",
    "/System/Library/Fonts/STHeiti Medium.ttc",
    "/System/Library/Fonts/Supplemental/Songti.ttc",
]

_CN_FONT_NAME = None
_CN_FONT_REGISTERED = False


def _register_cn_font() -> str:
    """
    注册中文字体，返回字体名称

    Returns:
        str: 可嵌入的中文字体名称；字体缺失时报告明确错误。
    """
    global _CN_FONT_NAME, _CN_FONT_REGISTERED
    if _CN_FONT_REGISTERED:
        return _CN_FONT_NAME

    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        for path in _CN_FONT_CANDIDATES:
            if not os.path.exists(path):
                continue
            try:
                font_name = f"CN_{os.path.basename(path).replace('.', '_')}"
                pdfmetrics.registerFont(TTFont(font_name, path))
                _CN_FONT_NAME = font_name
                break
            except Exception:
                continue
    except Exception:
        pass

    if _CN_FONT_NAME is None:
        raise RuntimeError("未找到可嵌入的中文字体，请检查应用 assets/fonts 是否完整")
    _CN_FONT_REGISTERED = True
    return _CN_FONT_NAME


def _kpi_rows(result: SimulationResult) -> List[List[str]]:
    """构造 KPI 表格行（中文标签 + 值）"""
    k = result.kpis
    duration_minutes = (result.elapsed_seconds if getattr(result, "run_id", "")
                        else result.duration_seconds) / 60.0
    unit = result.unit
    rows = [
        ["产线名称", result.line_name],
        ["实际经过时间(分钟)", f"{duration_minutes:.2f}"],
        ["计划仿真时长(分钟)", f"{result.duration_seconds / 60:.2f}"],
        [f"总产出({unit})", f"{result.total_output}"],
        [f"可发运产出({unit})", f"{k.get('shippable_quantity', result.total_output)}"],
        ["隔离批次", f"{k.get('rejected_batches', 0)}"],
        [f"瓶颈产能({unit}/h)", f"{k.get('bottleneck_capacity', 0):.1f}"],
        [f"预计日产量({unit})", f"{k.get('daily_output', 0):.0f}"],
        ["预计日人工成本(元)", f"{k.get('total_cost', 0):.2f}"],
        [f"预计单位人工成本(元/{unit})", f"{k.get('unit_cost', 0):.3f}"],
        ["产线平衡率", f"{k.get('balance_rate', 0) * 100:.1f}%"],
        [f"UPPH({unit}/人·h)", f"{k.get('upph', 0):.1f}"],
    ]
    rows.extend([
        ["运行ID", str(getattr(result, "run_id", ""))],
        ["配置修订", str(getattr(result, "config_revision", ""))],
        ["随机种子", str(getattr(result, "random_seed", None))],
        ["运行状态", str(getattr(result, "status", ""))],
        ["模型范围", ("配方批次流；不支持按工序人数反推烟油产能"
                       if getattr(result, "model_scope", "").startswith("recipe_batch")
                       else "离散工序流")],
        ["统计窗口(分钟)", f"{getattr(result, 'measurement_seconds', result.duration_seconds) / 60:.2f}"],
        ["预热(分钟)", f"{getattr(result, 'warmup_seconds', 0) / 60:.2f}"],
        ["本次人工成本(元)", f"{k.get('actual_labor_cost', k.get('labor_cost', 0)):.2f}"],
        [f"本次单位人工成本(元/{unit})",
         f"{k['actual_unit_cost']:.3f}" if k.get('actual_unit_cost') is not None else "无合格产出，无法计算"],
        ["成本边界", "仅人工；原料成本单列，不代表完整制造成本"],
    ])
    if k.get("material_cost") is not None:
        rows.append(["原料成本(元)", f"{k['material_cost']:.2f}"])
    if k.get("batch_cycle_min") is not None:
        rows.append(["批次周期(分钟)", f"{k['batch_cycle_min']:.1f}"])
    if k.get("batch_pass_rate") is not None:
        rows.append(["批次合格率", f"{k['batch_pass_rate']:.1%}"])
    if k.get("yield_rate") is not None:
        rows.append(["收率", f"{k['yield_rate']:.1%}"])
    if k.get("machine_oee") is not None:
        rows.append(["机台OEE", f"{k['machine_oee']:.1%}"])
    if k.get("cleaning_time_ratio") is not None:
        rows.append(["清洗时间占比", f"{k['cleaning_time_ratio']:.1%}"])
    if k.get("cost_per_liter") is not None:
        rows.append([f"人工元/{unit}（批次收率口径）", f"{k['cost_per_liter']:.3f}"])
    if k.get("cost_per_pouch") is not None:
        rows.append([f"人工元/{unit}（本次良品口径）", f"{k['cost_per_pouch']:.3f}"])
    return rows


def export_excel(result: SimulationResult, file_path: str) -> bool:
    """
    导出 Excel 报告

    Args:
        result: 仿真结果
        file_path: 保存路径（.xlsx）

    Returns:
        bool: 是否成功
    """
    try:
        import pandas as pd

        kpi_df = pd.DataFrame(_kpi_rows(result), columns=["指标", "值"])

        station_rows = []
        for station in result.station_outputs:
            station_rows.append({
                "工序ID": station,
                f"产出({result.unit})": result.station_outputs.get(station, 0),
                f"WIP({result.unit})": result.station_wips.get(station, 0),
            })
        station_df = pd.DataFrame(station_rows)

        alert_df = pd.DataFrame([
            {
                "时间(分钟)": round(a.timestamp_minutes, 1),
                "级别": a.severity,
                "类型": a.alert_type,
                "消息": a.message,
                "建议": a.suggestion,
            }
            for a in result.alerts
        ])

        wip_df = pd.DataFrame(result.wip_samples)
        changeover_df = pd.DataFrame(result.changeover_events)
        batch_df = pd.DataFrame(result.batch_results)
        quality_df = pd.DataFrame(result.quality_results)
        metrics_df = pd.DataFrame([
            {
                '工序ID': sid,
                '工序名': m.get('name', ''),
                '运行秒': m.get('running_sec', 0),
                '等待秒': m.get('waiting_sec', 0),
                '堵塞秒': m.get('blocked_sec', 0),
                '实际利用率': m.get('utilization', 0),
                'OEE可用率': m.get('oee_availability', 0),
                'OEE性能率': m.get('oee_performance', 0),
                'OEE合格率': m.get('oee_quality', 0),
                'OEE': m.get('oee_total', 0),
            }
            for sid, m in result.station_metrics.items()
        ])

        # P2 失衡分析
        imbalance_rows = []
        station_ids = list(result.station_metrics.keys())
        for i, sid in enumerate(station_ids):
            m = result.station_metrics[sid]
            ratio = None
            if i > 0:
                prev_cap = result.station_metrics[station_ids[i - 1]].get('capacity', 0)
                capacity = m.get('capacity', 0)
                ratio = round(capacity / prev_cap, 2) if prev_cap else None
            conclusion = []
            if m.get('utilization', 0) >= 0.95:
                conclusion.append('瓶颈')
            if m.get('waiting_sec', 0) >= 300:
                conclusion.append('饥饿')
            if m.get('blocked_sec', 0) >= 300:
                conclusion.append('堵塞')
            if ratio is not None and ratio < 0.8:
                conclusion.append('产能冗余')
            imbalance_rows.append({
                '工序ID': sid,
                '工序名': m.get('name', ''),
                '理论产能': m.get('capacity', 0),
                '上下游产能比': ratio if ratio is not None else '',
                '实际利用率': m.get('utilization', 0),
                '等待秒': m.get('waiting_sec', 0),
                '堵塞秒': m.get('blocked_sec', 0),
                '结论': '/'.join(conclusion) or '正常',
            })
        imbalance_df = pd.DataFrame(imbalance_rows)

        # V3.2 原料事件
        material_df = pd.DataFrame(result.material_events)

        # V3.3.1 产线配置
        lc = result.line_config or {}
        config_rows = [
            ["产线名称", lc.get("name", "")],
            ["生产类型", lc.get("production_type", "")],
            ["班次时长(小时)", lc.get("shift_hours", "")],
            ["休息(分钟)", lc.get("break_minutes", "")],
            ["时薪(元/h)", lc.get("worker_hourly_wage", "")],
            ["周期CIP(批次数)", lc.get("cip_interval_batches", 0)],
            ["周期CIP(小时)", lc.get("cip_interval_hours", 0)],
            [
                "换型矩阵",
                json.dumps(
                    lc.get("changeover_matrix") or {},
                    ensure_ascii=False,
                ),
            ],
        ]
        for key in ("recipes", "batches", "tanks", "labor_config", "cleanroom_limits", "skill_matrix"):
            config_rows.append([key, json.dumps(lc.get(key, []), ensure_ascii=False)])
        for material in lc.get("materials", []):
            config_rows.append([
                f"原料-{material.get('name', '')}",
                (
                    f"{material.get('initial_stock', 0)} {material.get('unit', '')}"
                    f"（单价 {material.get('unit_cost', 0)}）"
                ),
            ])
        for arrival in lc.get("material_arrivals", []):
            config_rows.append([
                f"到货-{arrival.get('material', '')}",
                f"{arrival.get('time_minutes', 0)}min × {arrival.get('quantity', 0)}",
            ])
        for station in lc.get("stations", []):
            config_rows.append([
                f"工序-{station.get('name', '')}",
                (
                    f"人数{station.get('worker_count', 0)}/机台{station.get('machine_count', 1)}/"
                    f"缓冲{station.get('buffer_capacity', 0)}/"
                    f"节拍{station.get('machine_takt') or '-'}/"
                    f"工种{station.get('job_role', '')}"
                ),
            ])
        config_df = pd.DataFrame(config_rows, columns=["配置项", "值"])

        directory = os.path.dirname(file_path)
        if directory and not os.path.exists(directory):
            os.makedirs(directory, exist_ok=True)

        with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
            kpi_df.to_excel(writer, sheet_name="KPI", index=False)
            station_df.to_excel(writer, sheet_name="工序产出", index=False)
            alert_df.to_excel(writer, sheet_name="报警记录", index=False)
            wip_df.to_excel(writer, sheet_name="WIP采样", index=False)
            changeover_df.to_excel(writer, sheet_name="切换事件", index=False)
            if not batch_df.empty:
                batch_df.to_excel(writer, sheet_name="批次结果", index=False)
            if not quality_df.empty:
                quality_df.to_excel(writer, sheet_name="质量门", index=False)
            if not metrics_df.empty:
                metrics_df.to_excel(writer, sheet_name="工序指标", index=False)
            if not imbalance_df.empty:
                imbalance_df.to_excel(writer, sheet_name="失衡分析", index=False)
            if not material_df.empty:
                material_df.to_excel(writer, sheet_name="原料事件", index=False)
            if not config_df.empty:
                config_df.to_excel(writer, sheet_name="产线配置", index=False)
            if getattr(result, "closing_tanks", []):
                pd.DataFrame(result.closing_tanks).to_excel(
                    writer, sheet_name="期末罐存", index=False
                )
            if getattr(result, "incomplete_batches", []):
                pd.DataFrame(result.incomplete_batches).to_excel(
                    writer, sheet_name="未完成批次", index=False
                )
            if getattr(result, "liquid_balance", {}):
                pd.DataFrame(list(result.liquid_balance.items()), columns=["指标", "数量"]).to_excel(
                    writer, sheet_name="液体守恒", index=False
                )

        return True
    except Exception:
        return False


def export_pdf(result: SimulationResult, file_path: str) -> bool:
    """
    导出 PDF 摘要报告

    Args:
        result: 仿真结果
        file_path: 保存路径（.pdf）

    Returns:
        bool: 是否成功
    """
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )

        font_name = _register_cn_font()

        directory = os.path.dirname(file_path)
        if directory and not os.path.exists(directory):
            os.makedirs(directory, exist_ok=True)

        doc = SimpleDocTemplate(
            file_path,
            pagesize=A4,
            rightMargin=20 * mm,
            leftMargin=20 * mm,
            topMargin=20 * mm,
            bottomMargin=20 * mm,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "CNTitle",
            parent=styles["Title"],
            fontName=font_name,
            fontSize=18,
            leading=24,
        )
        body_style = ParagraphStyle(
            "CNBody",
            parent=styles["Normal"],
            fontName=font_name,
            fontSize=10,
            leading=14,
        )
        header_style = ParagraphStyle(
            "CNHeader",
            parent=styles["Normal"],
            fontName=font_name,
            fontSize=11,
            leading=14,
        )

        def _pdf_table(rows, widths, header=True):
            wrapped = [
                [Paragraph(escape(str(value)).replace("\n", "<br/>"), body_style)
                 for value in row]
                for row in rows
            ]
            table = Table(wrapped, colWidths=widths, repeatRows=1 if header else 0)
            style = [
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
            if header:
                style.append(
                    ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey)
                )
            table.setStyle(TableStyle(style))
            return table

        story = [
            Paragraph(f"产线仿真报告 - {escape(result.line_name)}", title_style),
            Spacer(1, 8 * mm),
            Paragraph(
                f"实际经过时间：{(result.elapsed_seconds if getattr(result, 'run_id', '') else result.duration_seconds) / 60:.2f} 分钟　"
                f"总产出：{result.total_output} {result.unit}",
                body_style,
            ),
            Spacer(1, 6 * mm),
            Paragraph("一、核心 KPI", header_style),
            Spacer(1, 3 * mm),
        ]
        story.append(_pdf_table(_kpi_rows(result), [80 * mm, 90 * mm]))

        # 二、产线配置
        lc = result.line_config or {}
        config_rows = [["配置项", "值"]]
        config_rows += [
            ["产线名称", str(lc.get("name", ""))],
            ["生产类型", str(lc.get("production_type", ""))],
            ["班次时长(小时)", str(lc.get("shift_hours", ""))],
            ["休息(分钟)", str(lc.get("break_minutes", ""))],
            ["时薪(元/h)", str(lc.get("worker_hourly_wage", ""))],
            [
                "周期CIP",
                f"{lc.get('cip_interval_batches', 0)}批 / "
                f"{lc.get('cip_interval_hours', 0)}小时",
            ],
            [
                "换型矩阵",
                json.dumps(
                    lc.get("changeover_matrix") or {},
                    ensure_ascii=False,
                ),
            ],
        ]
        for material in lc.get("materials", []):
            config_rows.append([
                f"原料-{material.get('name', '')}",
                f"{material.get('initial_stock', 0)} {material.get('unit', '')}",
            ])
        for station in lc.get("stations", []):
            config_rows.append([
                f"工序-{station.get('name', '')}",
                (
                    f"人数{station.get('worker_count', 0)}/"
                    f"缓冲{station.get('buffer_capacity', 0)}/"
                    f"节拍{station.get('machine_takt') or '-'}"
                ),
            ])
        story.append(Spacer(1, 6 * mm))
        story.append(Paragraph("二、产线配置", header_style))
        story.append(Spacer(1, 3 * mm))
        story.append(_pdf_table(config_rows, [80 * mm, 90 * mm]))

        # 三、工序指标（含 OEE）
        if result.station_metrics:
            metric_rows = [["工序", "运行s", "等待s", "堵塞s", "利用率", "OEE"]]
            for sid, m in result.station_metrics.items():
                metric_rows.append([
                    str(m.get("name", sid)),
                    f"{m.get('running_sec', 0):.0f}",
                    f"{m.get('waiting_sec', 0):.0f}",
                    f"{m.get('blocked_sec', 0):.0f}",
                    f"{m.get('utilization', 0) * 100:.1f}%",
                    f"{m.get('oee_total', 0) * 100:.1f}%",
                ])
            story.append(Spacer(1, 6 * mm))
            story.append(Paragraph("三、工序指标（含 OEE）", header_style))
            story.append(Spacer(1, 3 * mm))
            story.append(_pdf_table(metric_rows, [35 * mm, 25 * mm, 25 * mm, 25 * mm, 30 * mm, 30 * mm]))

        # 四、失衡分析
        imbalance_rows = [["工序", "结论"]]
        station_ids = list(result.station_metrics.keys())
        for i, sid in enumerate(station_ids):
            m = result.station_metrics[sid]
            ratio = None
            if i > 0:
                prev_cap = result.station_metrics[station_ids[i - 1]].get("capacity", 0)
                capacity = m.get("capacity", 0)
                ratio = round(capacity / prev_cap, 2) if prev_cap else None
            conclusion = []
            if m.get("utilization", 0) >= 0.95:
                conclusion.append("瓶颈")
            if m.get("waiting_sec", 0) >= 300:
                conclusion.append("饥饿")
            if m.get("blocked_sec", 0) >= 300:
                conclusion.append("堵塞")
            if ratio is not None and ratio < 0.8:
                conclusion.append("产能冗余")
            imbalance_rows.append([
                str(m.get("name", sid)),
                "/".join(conclusion) or "正常",
            ])
        story.append(Spacer(1, 6 * mm))
        story.append(Paragraph("四、失衡分析", header_style))
        story.append(Spacer(1, 3 * mm))
        story.append(_pdf_table(imbalance_rows, [80 * mm, 90 * mm]))

        if result.alerts:
            story.append(Spacer(1, 6 * mm))
            story.append(Paragraph("五、报警记录", header_style))
            story.append(Spacer(1, 3 * mm))
            alert_rows = [["时间(分钟)", "级别", "消息"]]
            for alert in result.alerts[:50]:
                alert_rows.append([
                    f"{alert.timestamp_minutes:.1f}",
                    alert.severity,
                    alert.message,
                ])
            story.append(_pdf_table(alert_rows, [25 * mm, 25 * mm, 120 * mm]))

        if result.batch_results:
            story.append(Spacer(1, 6 * mm))
            story.append(Paragraph("六、批次结果", header_style))
            story.append(Spacer(1, 3 * mm))
            batch_rows = [["批次ID", "配方", "周期(分钟)", "合格率", f"产出({result.unit})"]]
            for b in result.batch_results:
                batch_rows.append([
                    str(b.get("batch_id", "")),
                    str(b.get("recipe_name", "")),
                    f"{b.get('cycle_min', 0):.1f}",
                    f"{b.get('pass_rate', 0) * 100:.1f}%",
                    f"{b.get('yield_l', 0):.1f}",
                ])
            story.append(_pdf_table(batch_rows, [30 * mm, 35 * mm, 35 * mm, 30 * mm, 40 * mm]))

        if result.changeover_events:
            story.append(Spacer(1, 6 * mm))
            title = "换型事件" if result.batch_results else "六、换型事件"
            story.append(Paragraph(title, header_style))
            story.append(Spacer(1, 3 * mm))
            event_rows = [["工序", "申请(分)", "开始(分)", "结束(分)", "时长(分)", "状态"]]
            for event in result.changeover_events:
                times = [event.get(key) for key in ('requested_time', 'start_time', 'end_time')]
                event_rows.append([
                    str(event.get('station_name', event.get('station_id', ''))),
                    *[f"{value / 60:.2f}" if value is not None else "-" for value in times],
                    str(event.get('minutes', '')),
                    str(event.get('status', '')),
                ])
            story.append(_pdf_table(event_rows, [40 * mm] + [26 * mm] * 5))

        if result.material_events:
            story.append(Spacer(1, 6 * mm))
            story.append(Paragraph("七、原料事件", header_style))
            story.append(Spacer(1, 3 * mm))
            material_rows = [["时间", "类型", "原料", "数量"]]
            for event in result.material_events[:50]:
                material_rows.append([
                    f"{event.get('time', 0):.0f}s",
                    str(event.get("type", "")),
                    str(event.get("material", "")),
                    str(event.get("quantity", "")),
                ])
            story.append(_pdf_table(material_rows, [35 * mm, 35 * mm, 50 * mm, 50 * mm]))

        story.append(Spacer(1, 6 * mm))
        if getattr(result, "liquid_balance", {}):
            story.append(Paragraph("八、全程液体守恒与未完成批次", header_style))
            story.append(Spacer(1, 3 * mm))
            balance_labels = {
                "initial_tank_quantity": "期初罐存", "produced_good_quantity": "完成QC良品量",
                "drained_quantity": "罐已排出量", "closing_tank_quantity": "期末罐存",
                "pending_tank_quantity": "待入罐量",
            }
            rows = [["指标（覆盖整次运行，含预热）", f"数量({result.unit})"]]
            rows.extend([[balance_labels.get(key, key), f"{value:.3f}"]
                         for key, value in result.liquid_balance.items()])
            rows.append(["未完成批次数", str(len(result.incomplete_batches))])
            story.append(_pdf_table(rows, [110 * mm, 60 * mm]))
            for batch in result.incomplete_batches:
                story.append(Paragraph(
                    escape(f"批次{batch.get('id', '')}：{batch.get('stage', batch.get('status', ''))}；"
                           f"已入罐{batch.get('stored_quantity', 0)}，"
                           f"待入罐{batch.get('pending_tank_quantity', 0)} {result.unit}；未计入完成放行量"),
                    body_style,
                ))
            story.append(Spacer(1, 6 * mm))
        story.append(Paragraph(
            "注：本报告由仿真工具自动生成，结果作为产线设计参考，"
            "建议结合实际生产数据校准 OEE 等参数。",
            body_style,
        ))

        _keep_section_headers(story)
        doc.build(story)
        return True
    except Exception:
        return False


def export_hr_report(summary: dict, file_path: str) -> bool:
    """导出人力规划 Excel 报告（V3.3 迭代）"""
    try:
        import pandas as pd
        costs = summary.get("costs", {})

        overview_df = pd.DataFrame([
            ["目标日产量", summary.get("daily_target", 0)],
            ["每日有效工时(小时)", summary.get("effective_hours_per_day", 0)],
            ["每小时需求", summary.get("required_hourly", 0)],
            ["每班最低岗位人数", summary.get("headcount_per_shift", 0)],
            ["每日班组数", summary.get("shifts_per_day", 1)],
            ["每班覆盖编制(含缺勤)", summary.get("covered_headcount_per_shift", 0)],
            ["全天覆盖编制(含缺勤)", summary.get("covered_headcount", 0)],
            ["全天编制合计(不含缺勤覆盖)", summary.get("total_headcount", 0)],
            ["产能可行", "是" if summary.get("capacity_feasible", True) else "否"],
            ["人力计划可达", "是" if summary.get("feasible", True) else "否"],
            ["约束说明", "；".join(summary.get("constraints", []))],
            ["当前在岗合计", summary.get("current_total", 0)],
            ["富余人数(在岗−覆盖编制)", summary.get("excess_headcount", 0)],
            ["一次性招聘建议", summary.get("one_time_hiring", 0)],
            ["初始缺口(覆盖编制−在岗)", summary.get("initial_gap", 0)],
            ["覆盖编制达产天数", summary.get("days_to_full") if summary.get("days_to_full") is not None
             else "计划内无法达产"],
        ], columns=["指标", "值"])

        derivation_df = pd.DataFrame([
            {
                "工序": d.get("name", ""),
                "是否瓶颈": "是" if d.get("is_bottleneck") else "",
                "配置人数": d.get("configured", 0),
                "单件耗时(秒)": d.get("process_time", 0),
                "机台节拍(秒)": d.get("machine_takt") or "",
                "OEE": d.get("oee", 0),
                "效率": d.get("efficiency", 0),
                "切换+清洗(分钟)": d.get("changeover_minutes", 0),
                "协作模式": d.get("collaboration", ""),
                "工序产能(/h)": d.get("capacity", 0),
                "单人产能(/h)": d.get("per_head", 0),
                "需求人数": d.get("needed", 0),
            }
            for d in summary.get("station_derivation", [])
        ])

        role_df = pd.DataFrame([
            {"工种": role, "每班人数": summary.get("headcount_by_role_per_shift", {}).get(role, n),
             "最低全天编制": n,
             "覆盖全天编制": summary.get("covered_headcount_by_role", {}).get(role, n)}
            for role, n in summary.get("headcount_by_role", {}).items()
        ] + [{
            "工种": "合计", "每班人数": summary.get("headcount_per_shift", 0),
            "最低全天编制": summary.get("total_headcount", 0),
            "覆盖全天编制": summary.get("covered_headcount", 0),
        }])

        current_df = pd.DataFrame([
            {"工种": role, "在岗人数": n}
            for role, n in summary.get("current_headcount", {}).items()
        ] + [{
            "工种": "合计", "在岗人数": summary.get("current_total", 0),
        }])

        cost_df = pd.DataFrame([
            ["总人数", costs.get("headcount", 0)],
            ["在岗覆盖人数（含缺勤）", costs.get("covered_headcount", 0)],
            ["日人力成本(元)", costs.get("daily_labor_cost", 0)],
            ["月人力成本(元)", costs.get("monthly_labor_cost", 0)],
            ["月招聘/培训(元)", costs.get("monthly_recruit_training", 0)],
            ["月总成本(元)", costs.get("monthly_total", 0)],
            ["单位人力成本(元/单位)", costs.get("per_unit_labor_cost", 0)],
        ], columns=["指标", "值"])

        gap_rows = []
        for row in summary.get("weekly_gap", []):
            gap_rows.append({
                "周": row.get("week", 0),
                "周底还缺(累计)": row.get("total_gap", 0),
                "本周新增需求": row.get("new_gap", 0),
                **{f"缺口_{role}": v for role, v in row.get("gap", {}).items()},
            })
        gap_df = pd.DataFrame(gap_rows)

        directory = os.path.dirname(file_path)
        if directory and not os.path.exists(directory):
            os.makedirs(directory, exist_ok=True)

        with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
            overview_df.to_excel(writer, sheet_name="概览", index=False)
            derivation_df.to_excel(writer, sheet_name="计算明细", index=False)
            role_df.to_excel(writer, sheet_name="人力需求", index=False)
            current_df.to_excel(writer, sheet_name="当前在岗", index=False)
            cost_df.to_excel(writer, sheet_name="成本", index=False)
            gap_df.to_excel(writer, sheet_name="招聘缺口", index=False)

        from openpyxl import load_workbook
        from openpyxl.styles import Font

        workbook = load_workbook(file_path)
        header_font = Font(bold=True, color="FFFFFF")
        for sheet_name, widths in (
            ("概览", [36, 20]),
            ("计算明细", [16, 9, 9, 12, 12, 8, 8, 13, 10, 12, 12, 9]),
            ("人力需求", [24, 12]),
            ("当前在岗", [24, 12]),
            ("成本", [30, 18]),
            ("招聘缺口", [10, 14, 14]),
        ):
            worksheet = workbook[sheet_name]
            for cell in worksheet[1]:
                cell.font = header_font
            for idx, width in enumerate(widths, 1):
                worksheet.column_dimensions[
                    worksheet.cell(row=1, column=idx).column_letter
                ].width = width
            worksheet.freeze_panes = "A2"
        workbook.save(file_path)
        return True
    except Exception:
        return False


def export_hr_pdf(summary: dict, file_path: str) -> bool:
    """导出人力规划 PDF 报告（V3.3 迭代）"""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )

        font_name = _register_cn_font()
        directory = os.path.dirname(file_path)
        if directory and not os.path.exists(directory):
            os.makedirs(directory, exist_ok=True)

        doc = SimpleDocTemplate(
            file_path, pagesize=A4,
            rightMargin=20 * mm, leftMargin=20 * mm,
            topMargin=20 * mm, bottomMargin=20 * mm,
        )
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "CNTitle", parent=styles["Title"],
            fontName=font_name, fontSize=18, leading=24,
        )
        body_style = ParagraphStyle(
            "CNBody", parent=styles["Normal"],
            fontName=font_name, fontSize=10, leading=14,
        )

        costs = summary.get("costs", {})

        def _table(rows, widths, header=True):
            header_style = ParagraphStyle("HRHeaderCell", parent=body_style, textColor=colors.white)
            wrapped = [
                [Paragraph(escape(str(value)).replace("\n", "<br/>"),
                           header_style if header and i == 0 else body_style)
                 for value in row]
                for i, row in enumerate(rows)
            ]
            table = Table(wrapped, colWidths=widths, repeatRows=1 if header else 0)
            style = [
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
            if header:
                style += [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2329")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ]
            table.setStyle(TableStyle(style))
            return table

        story = [
            Paragraph("人力规划报告", title_style),
            Spacer(1, 6 * mm),
            Paragraph("一、概览", body_style),
            Spacer(1, 2 * mm),
        ]
        overview_rows = [
            ["指标", "值"],
            ["目标日产量", summary.get("daily_target", 0)],
            ["每日有效工时(小时)", summary.get("effective_hours_per_day", 0)],
            ["每小时需求", summary.get("required_hourly", 0)],
            ["每班最低岗位人数", summary.get("headcount_per_shift", 0)],
            ["每日班组数", summary.get("shifts_per_day", 1)],
            ["每班覆盖编制(含缺勤)", summary.get("covered_headcount_per_shift", 0)],
            ["全天覆盖编制(含缺勤)", summary.get("covered_headcount", 0)],
            ["全天编制合计(不含缺勤覆盖)", summary.get("total_headcount", 0)],
            ["产能可行", "是" if summary.get("capacity_feasible", True) else "否"],
            ["人力计划可达", "是" if summary.get("feasible", True) else "否"],
            ["约束说明", "；".join(summary.get("constraints", []))],
            ["当前在岗合计", summary.get("current_total", 0)],
            ["富余人数(在岗−覆盖编制)", summary.get("excess_headcount", 0)],
            ["一次性招聘建议", summary.get("one_time_hiring", 0)],
            ["覆盖编制达产天数", summary.get("days_to_full") if summary.get("days_to_full") is not None
             else "计划内无法达产"],
        ]
        story.append(_table(overview_rows, [110 * mm, 50 * mm]))
        story.append(Spacer(1, 4 * mm))
        story.append(Paragraph("二、计算明细", body_style))
        story.append(Spacer(1, 2 * mm))
        derivation_rows = [["工序", "配置", "单人产能", "需求", "备注"]]
        derivation_rows += [
            [
                d.get("name", ""),
                str(d.get("configured", 0)),
                f"{d.get('per_head', 0)}/h",
                str(d.get("needed", 0)),
                "瓶颈" if d.get("is_bottleneck") else "",
            ]
            for d in summary.get("station_derivation", [])
        ]
        story.append(_table(
            derivation_rows, [55 * mm, 20 * mm, 45 * mm, 20 * mm, 30 * mm]
        ))
        story.append(Spacer(1, 4 * mm))
        story.append(Paragraph("三、人力需求（按工种）", body_style))
        story.append(Spacer(1, 2 * mm))
        role_rows = [["工种", "最低全天编制", "覆盖全天编制"]]
        role_rows += [
            [role, str(count), str(summary.get('covered_headcount_by_role', {}).get(role, count))]
            for role, count in summary.get("headcount_by_role", {}).items()
        ]
        role_rows.append(["合计", str(summary.get("total_headcount", 0)),
                          str(summary.get('covered_headcount', 0))])
        story.append(_table(role_rows, [70 * mm, 50 * mm, 50 * mm]))
        story.append(Spacer(1, 4 * mm))
        story.append(Paragraph("四、成本测算", body_style))
        story.append(Spacer(1, 2 * mm))
        cost_rows = [
            ["指标", "值"],
            ["总人数", costs.get("headcount", 0)],
            ["在岗覆盖人数（含缺勤）", costs.get("covered_headcount", 0)],
            ["日人力成本(元)", costs.get("daily_labor_cost", 0)],
            ["月总成本(元)", costs.get("monthly_total", 0)],
            ["单位人力成本(元/单位)", costs.get("per_unit_labor_cost", 0)],
        ]
        story.append(_table(cost_rows, [80 * mm, 40 * mm]))
        story.append(Spacer(1, 4 * mm))
        story.append(Paragraph("五、招聘缺口（前 12 周）", body_style))
        story.append(Spacer(1, 2 * mm))
        gap_rows = [["周", "周底还缺(累计)", "本周新增需求"]]
        gap_rows += [
            [
                str(row.get("week", 0)),
                str(row.get("total_gap", 0)),
                str(row.get("new_gap", 0)),
            ]
            for row in summary.get("weekly_gap", [])[:12]
        ]
        story.append(_table(gap_rows, [30 * mm, 70 * mm, 70 * mm]))
        story.append(Spacer(1, 4 * mm))
        story.append(Paragraph(
            f"六、结论：一次性招聘建议 {summary.get('one_time_hiring', 0)} 人；"
            + (f"预计 {summary['days_to_full']} 天满足覆盖编制。"
               if summary.get('days_to_full') is not None else "当前招聘/培训计划内无法达产。"),
            body_style,
        ))
        _keep_section_headers(story)
        doc.build(story)
        return True
    except Exception:
        return False


def export_report(result: SimulationResult, file_path: str) -> bool:
    """
    按文件扩展名导出报告

    Args:
        result: 仿真结果
        file_path: 保存路径（.xlsx 或 .pdf）

    Returns:
        bool: 是否成功
    """
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".xlsx":
        return export_excel(result, file_path)
    if ext == ".pdf":
        return export_pdf(result, file_path)
    return False
