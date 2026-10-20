"""跨行业边界：兼容、计量、可行节拍与成品守恒。"""
import copy

import pytest
from openpyxl import load_workbook
from src.gui_panels import build_template_line
from src.history import build_snapshot
from src.models import JobRole, ProductionLine, ProductionType, Station
from src.optimizer import optimize
from src.reporting import export_excel, export_pdf
from src.simulation import SimulationEngine
from src.templates import build_process_template
from src.utils import create_excel_template, import_from_excel, validate_production_line


def hardware_line():
    line = ProductionLine('合成硬件测试线', shift_hours=1, break_minutes=0, output_unit='台')
    line.stations = [Station('s01', '功能测试', 10, 1, machine_takt=10,
                             oee=1, efficiency=1, job_role=JobRole.TEST_OPERATOR)]
    return line


def test_legacy_json_keeps_unit_downtime_and_input_unchanged():
    old = {'name': '旧方案', 'stations': [dict(id='s01', name='组装',
                                             process_time=10, worker_count=1)]}
    original = copy.deepcopy(old)
    line = ProductionLine.from_dict(old)
    assert line.get_unit() == '颗'
    assert line.stations[0].changeover_time == 45
    assert old == original
    assert ProductionLine.from_dict(line.to_dict()).to_dict() == line.to_dict()
    assert ProductionLine('新方案').get_unit() == '件'
    assert Station('s01', '装配', 10, 1).changeover_time == 0


@pytest.mark.parametrize('role', [JobRole.ASSEMBLER, JobRole.TEST_OPERATOR, JobRole.MAINTENANCE])
def test_hardware_roles_and_unit_roundtrip(role):
    line = hardware_line()
    line.stations[0].job_role = role
    restored = ProductionLine.from_dict(line.to_dict())
    assert restored.get_unit() == '台'
    assert restored.stations[0].job_role == role
    assert validate_production_line(restored)[0]


@pytest.mark.parametrize('unit', ['件', '台', '套'])
def test_unit_changes_label_without_changing_output(unit, tmp_path):
    line = hardware_line()
    line.output_unit = unit
    result = SimulationEngine(line).run_sync(60 / 3600)
    assert result.total_output == 5
    assert result.unit == unit
    assert build_snapshot(line, result)['unit'] == unit
    path = tmp_path / 'report.xlsx'
    export_excel(result, str(path))
    wb = load_workbook(path)
    assert any(f'总产出({unit})' == cell.value
               for row in wb['KPI'] for cell in row)
    pdf = tmp_path / 'report.pdf'
    export_pdf(result, str(pdf))
    assert pdf.read_bytes().startswith(b'%PDF')


@pytest.mark.parametrize('minimum', [None, 8.0])
def test_optimizer_respects_physical_lower_bound(minimum):
    line = hardware_line()
    line.stations[0].min_machine_takt = minimum
    results = optimize(line, objective='output', generations=2, population=4,
                       duration_hours=60 / 3600)
    for row in results:
        assert row['params']['s01']['machine_takt'] >= (minimum or 10)


@pytest.mark.parametrize('minimum', [0, -1, 11, float('nan'), float('inf')])
def test_invalid_minimum_is_rejected_before_search(minimum):
    line = hardware_line()
    line.stations[0].min_machine_takt = minimum
    assert not validate_production_line(line)[0]
    with pytest.raises(ValueError):
        optimize(line, generations=1, population=1)


def test_excel_unit_constraint_and_legacy_fallback(tmp_path):
    path = str(tmp_path / 'hardware.xlsx')
    assert create_excel_template(path)
    wb = load_workbook(path)
    ws = wb['工序']
    ws.cell(2, 9, 10)
    ws.cell(2, 18, 8)
    wb['产线配置'].append(['output_unit', '台'])
    wb.save(path)
    line, error = import_from_excel(path)
    assert error is None
    assert line.get_unit() == '台'
    assert line.stations[0].min_machine_takt == 8
    ws.cell(2, 18, 11)
    wb.save(path)
    assert import_from_excel(path)[0] is None
    ws.cell(2, 18).value = None
    config = wb['产线配置']
    for row in range(config.max_row, 1, -1):
        if config.cell(row, 1).value == 'output_unit':
            config.delete_rows(row)
    wb.save(path)
    assert import_from_excel(path)[0].get_unit() == '颗'


@pytest.mark.parametrize('kind', ['simple', 'standard', 'complex'])
def test_generic_assembly_templates_run_without_tobacco(kind):
    line = build_template_line(kind)
    assert not any(word in str(line.to_dict()) for word in ('尼古丁', '烟油', '棉芯', '注油'))
    assert SimulationEngine(line).run_sync(1).total_output > 0


@pytest.mark.parametrize('kind', ['liquid_filling', 'pouch_packaging'])
def test_generic_process_templates_run_without_tobacco(kind):
    line = build_process_template(kind)
    assert not any(word in str(line.to_dict()) for word in ('尼古丁', '烟草', '薄荷'))
    assert SimulationEngine(line).run_sync(8).total_output > 0


@pytest.mark.parametrize('unit', ['a b', '\n', 'a'*13])
def test_invalid_units_rejected(unit):
    line = hardware_line()
    line.output_unit = unit
    assert not validate_production_line(line)[0]


def test_liquid_mass_unit_cannot_be_relabelled_as_count():
    line = build_process_template('liquid_filling')
    line.output_unit = '台'
    assert not validate_production_line(line)[0]
