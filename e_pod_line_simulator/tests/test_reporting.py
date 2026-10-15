"""报告导出单元测试"""

import os

from src.models import ProductionLine, Station, create_liquid_line
from src.reporting import (
    export_excel,
    export_hr_pdf,
    export_hr_report,
    export_pdf,
    export_report,
)
from src.simulation import SimulationEngine


def make_result():
    line = ProductionLine("报告测试", shift_hours=1, break_minutes=0)
    line.add_station(Station("s01", "注油", 1.0, 1))
    line.add_station(Station("s02", "包装", 1.0, 1))
    return SimulationEngine(line).run_sync(duration_hours=0.1)


def test_export_excel(tmp_path):
    path = str(tmp_path / "report.xlsx")
    assert export_excel(make_result(), path) is True
    assert os.path.exists(path)

    from openpyxl import load_workbook

    wb = load_workbook(path)
    assert "KPI" in wb.sheetnames
    assert "工序产出" in wb.sheetnames
    assert "报警记录" in wb.sheetnames
    assert "WIP采样" in wb.sheetnames
    assert "切换事件" in wb.sheetnames
    assert "工序指标" in wb.sheetnames
    assert "产线配置" in wb.sheetnames
    assert "失衡分析" in wb.sheetnames


def test_export_pdf(tmp_path):
    path = str(tmp_path / "report.pdf")
    assert export_pdf(make_result(), path) is True
    assert os.path.exists(path)
    with open(path, "rb") as f:
        assert f.read(4) == b"%PDF"


def test_export_report_dispatch(tmp_path):
    xlsx_path = str(tmp_path / "a.xlsx")
    pdf_path = str(tmp_path / "b.pdf")
    bad_path = str(tmp_path / "c.txt")
    result = make_result()
    assert export_report(result, xlsx_path) is True
    assert export_report(result, pdf_path) is True
    assert export_report(result, bad_path) is False


def test_liquid_report_contains_batch_sheets(tmp_path):
    line = create_liquid_line()
    result = SimulationEngine(line).run_sync(duration_hours=24.0)
    path = str(tmp_path / "liquid_report.xlsx")
    assert export_excel(result, path) is True

    from openpyxl import load_workbook

    wb = load_workbook(path)
    assert "批次结果" in wb.sheetnames
    assert "质量门" in wb.sheetnames
    assert "产线配置" in wb.sheetnames
    kpi_values = [
        str(c.value) for row in wb["KPI"].iter_rows() for c in row if c.value
    ]
    assert any("千克" in v for v in kpi_values)
    assert any("原料成本" in v for v in kpi_values)


def test_export_hr_report_excel_and_pdf(tmp_path):
    from src.hr_planning import (
        LaborCostConfig,
        LearningCurveConfig,
        ShiftPlan,
        build_hr_summary,
    )

    line = ProductionLine("HR报告测试", shift_hours=8, break_minutes=60)
    line.add_station(Station(
        "s01", "组装", 60.0, 2,
        oee=1.0, efficiency=1.0, changeover_time=0,
    ))
    summary = build_hr_summary(
        line,
        840,
        ShiftPlan(shifts_per_day=1, shift_hours=8, break_minutes=60),
        LaborCostConfig(),
        LearningCurveConfig(ramp_days=30),
        current={"general": 2},
        hiring_plan=[],
    )

    xlsx = str(tmp_path / "hr.xlsx")
    pdf = str(tmp_path / "hr.pdf")
    assert export_hr_report(summary, xlsx) is True
    assert export_hr_pdf(summary, pdf) is True

    from openpyxl import load_workbook

    wb = load_workbook(xlsx)
    assert "概览" in wb.sheetnames
    assert "计算明细" in wb.sheetnames
    assert "人力需求" in wb.sheetnames
    assert "当前在岗" in wb.sheetnames
    assert "成本" in wb.sheetnames
    assert "招聘缺口" in wb.sheetnames
    overview_values = [
        str(c.value) for row in wb["概览"].iter_rows() for c in row if c.value
    ]
    assert any("每小时需求" in v for v in overview_values)


def test_partial_liquid_report_keeps_run_identity_and_mass_balance(tmp_path):
    """A tank-blocked batch must remain visible instead of becoming finished goods."""
    from openpyxl import load_workbook
    from src.models import Batch, ProductionType, Recipe, Tank

    line = ProductionLine("小罐守恒", production_type=ProductionType.LIQUID_FILLING)
    line.recipes = [Recipe(
        "R", batch_volume_l=100, yield_rate=1, mixing_time_min=0,
        aging_time_min=0, filling_rate_l_per_h=6000, qc_time_min=0, clean_time_min=0,
    )]
    line.batches = [Batch("B", "R", 100)]
    line.tanks = [Tank("T", "成品罐", 10)]
    result = SimulationEngine(line).run_sync(70 / 3600)
    path = tmp_path / "partial.xlsx"
    assert export_excel(result, str(path))
    wb = load_workbook(path, data_only=True)
    assert {"期末罐存", "未完成批次", "液体守恒"}.issubset(wb.sheetnames)
    values = [cell.value for row in wb["KPI"] for cell in row]
    assert result.run_id in values
    assert result.config_revision in values
    assert result.total_output == 0
    assert result.liquid_balance["closing_tank_quantity"] == 10
    assert result.liquid_balance["pending_tank_quantity"] == 90
    assert "B" in [cell.value for row in wb["未完成批次"] for cell in row]
    assert 90 in [cell.value for row in wb["液体守恒"] for cell in row]
    wb.close()


def test_bundled_chinese_font_is_embedded_in_pdf(tmp_path, monkeypatch):
    """Reports must carry their font, even on a machine without Chinese fonts."""
    import src.reporting as reporting

    monkeypatch.setattr(reporting, "_CN_FONT_CANDIDATES", [reporting._BUNDLED_CN_FONT])
    monkeypatch.setattr(reporting, "_CN_FONT_NAME", None)
    monkeypatch.setattr(reporting, "_CN_FONT_REGISTERED", False)
    path = tmp_path / "portable.pdf"
    assert export_pdf(make_result(), str(path))
    pdf = path.read_bytes()
    assert b"/FontFile2" in pdf
    assert b"/ToUnicode" in pdf
    assert "NotoSansSC" in reporting._CN_FONT_NAME


def test_pdf_keeps_changeover_request_actual_start_and_end(tmp_path, monkeypatch):
    from reportlab.platypus import SimpleDocTemplate, Table

    result = make_result()
    result.changeover_events = [{
        'station_name': '注油', 'requested_time': 90, 'start_time': 120,
        'end_time': 180, 'minutes': 1, 'status': 'completed',
    }]
    table_text = []
    original_build = SimpleDocTemplate.build

    def inspect_build(document, story, *args, **kwargs):
        for item in story:
            if isinstance(item, Table):
                table_text.extend(
                    cell.getPlainText() if hasattr(cell, 'getPlainText') else str(cell)
                    for row in item._cellvalues for cell in row
                )
        return original_build(document, story, *args, **kwargs)

    monkeypatch.setattr(SimpleDocTemplate, 'build', inspect_build)
    assert export_pdf(result, str(tmp_path / 'changeover.pdf'))
    for expected in ('申请(分)', '开始(分)', '结束(分)', '1.50', '2.00', '3.00', 'completed'):
        assert expected in table_text
